"""Audit the cross-year IPEDS schema needed for the analysis.

This script verifies that required variables and aggregation codes are present
in every downloaded year before analysis begins. It writes a compact tracked
CSV so the schema review is visible without committing the raw source files.
"""

from __future__ import annotations

import csv
from pathlib import Path
import sys

from openpyxl import load_workbook


START_YEAR = 2015
END_YEAR = 2024
ROOT = Path(__file__).resolve().parents[1]
RAW_ROOT = ROOT / "data" / "raw" / "ipeds"
OUTPUT = ROOT / "data" / "schema_audit.csv"

EF_REQUIRED = ("UNITID", "EFALEVEL", "EFTOTLT", "EFNRALT")
HD_REQUIRED = ("UNITID", "INSTNM", "STABBR")
EXPECTED_LEVELS = {
    "1": "All students total",
    "2": "All students, Undergraduate total",
    "12": "All students, Graduate",
}


def workbook_path(year: int, table: str) -> Path:
    folder = RAW_ROOT / str(year) / f"{table}_Dict"
    matches = sorted(folder.glob("*.xlsx")) + sorted(folder.glob("*.xlsm"))
    if not matches:
        raise RuntimeError(f"No dictionary workbook found in {folder}")
    return matches[0]


def sheet_by_name(workbook, name: str):
    for sheet_name in workbook.sheetnames:
        if sheet_name.casefold() == name.casefold():
            return workbook[sheet_name]
    raise RuntimeError(f"Workbook has no {name!r} sheet: {workbook.sheetnames}")


def varlist(workbook) -> dict[str, str]:
    sheet = sheet_by_name(workbook, "varlist")
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    columns = {str(value).casefold(): index for index, value in enumerate(header)}
    name_index = columns["varname"]
    title_index = columns["vartitle"]
    result: dict[str, str] = {}
    for row in rows:
        if row[name_index] is not None:
            result[str(row[name_index])] = "" if row[title_index] is None else str(row[title_index])
    return result


def frequency_rows(workbook, variable: str) -> dict[str, str]:
    # Historical revised files commonly include FrequenciesRV. Prefer that
    # sheet when present because the analysis will prefer revised CSV files.
    names = {name.casefold(): name for name in workbook.sheetnames}
    selected = names.get("frequenciesrv") or names.get("frequencies")
    if selected is None:
        raise RuntimeError(f"No frequency sheet found: {workbook.sheetnames}")

    rows: dict[str, str] = {}
    for row in workbook[selected].iter_rows(values_only=True):
        if variable not in tuple("" if value is None else str(value) for value in row):
            continue

        values = ["" if value is None else str(value) for value in row]
        var_index = values.index(variable)
        trailing = values[var_index + 1 :]

        # The old and new dictionary formats place value/label in slightly
        # different columns. Find the first integer-like code after varName,
        # then use the following nonempty text value as its label.
        code = None
        label = None
        for index, value in enumerate(trailing):
            if (
                code is None
                and value.lstrip("-").isdigit()
                and abs(int(value)) <= 100
            ):
                code = value
                for candidate in trailing[index + 1 :]:
                    if candidate and not candidate.replace(",", "").replace(".", "").isdigit():
                        label = candidate
                        break
                if label is not None:
                    break
        if code is not None and label is not None:
            rows[code] = label
    return rows


def preferred_csv(year: int, table: str) -> str:
    folder = RAW_ROOT / str(year) / table
    revised = sorted(folder.glob("*_rv.csv"))
    if revised:
        return revised[0].name
    regular = sorted(folder.glob("*.csv"))
    if not regular:
        raise RuntimeError(f"No CSV found in {folder}")
    return regular[0].name


def main() -> int:
    records: list[dict[str, str | int]] = []
    failures: list[str] = []

    for year in range(START_YEAR, END_YEAR + 1):
        ef_table = f"EF{year}A"
        hd_table = f"HD{year}"

        ef_workbook = load_workbook(workbook_path(year, ef_table), read_only=True, data_only=True)
        hd_workbook = load_workbook(workbook_path(year, hd_table), read_only=True, data_only=True)
        ef_vars = varlist(ef_workbook)
        hd_vars = varlist(hd_workbook)
        levels = frequency_rows(ef_workbook, "EFALEVEL")

        missing_ef = [name for name in EF_REQUIRED if name not in ef_vars]
        missing_hd = [name for name in HD_REQUIRED if name not in hd_vars]
        bad_levels = [
            f"{code}={levels.get(code)!r}"
            for code, expected in EXPECTED_LEVELS.items()
            if levels.get(code) != expected
        ]

        status = "PASS"
        notes: list[str] = []
        if missing_ef:
            status = "FAIL"
            notes.append(f"missing EF vars: {', '.join(missing_ef)}")
        if missing_hd:
            status = "FAIL"
            notes.append(f"missing HD vars: {', '.join(missing_hd)}")
        if bad_levels:
            status = "FAIL"
            notes.append(f"unexpected EFALEVEL labels: {', '.join(bad_levels)}")

        if status == "FAIL":
            failures.append(f"{year}: {'; '.join(notes)}")

        records.append(
            {
                "year": year,
                "status": status,
                "enrollment_csv": preferred_csv(year, ef_table),
                "directory_csv": preferred_csv(year, hd_table),
                "nonresident_variable": "EFNRALT",
                "nonresident_title": ef_vars.get("EFNRALT", ""),
                "total_variable": "EFTOTLT",
                "level_all": levels.get("1", ""),
                "level_undergraduate": levels.get("2", ""),
                "level_graduate": levels.get("12", ""),
                "state_variable": "STABBR",
                "notes": "; ".join(notes),
            }
        )

    with OUTPUT.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(records[0].keys()))
        writer.writeheader()
        writer.writerows(records)

    print(f"Wrote {OUTPUT}")
    for record in records:
        print(
            f"{record['year']}: {record['status']} | "
            f"{record['enrollment_csv']} | {record['nonresident_title']}"
        )

    if failures:
        print("Schema audit failures:", file=sys.stderr)
        for failure in failures:
            print(f"- {failure}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

