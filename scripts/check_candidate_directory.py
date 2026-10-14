"""Capture a structural preflight for the available candidate-year IPEDS directory.

This is an online monitoring lane, not part of the reproducible offline build.
It is intentionally narrower than the annual schema audit: when the next-year
HD data and dictionary are available before Fall Enrollment, it checks the
directory-side fields and institution-universe continuity that can be reviewed
without pretending the analysis year is ready to advance.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from project_config import CORE_STATES, END_YEAR


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
RAW_ROOT = DATA / "raw" / "ipeds"
RELEASE_SNAPSHOT = DATA / "release_monitor_snapshot.json"
DEFAULT_OUTPUT = DATA / "candidate_directory_preflight.json"
DEFAULT_CHANGE_REVIEW_OUTPUT = DATA / "candidate_directory_change_review.csv"
USER_AGENT = "Mozilla/5.0 (compatible; ipeds-portfolio-directory-preflight/1.0)"

REQUIRED_FIELDS = ("UNITID", "INSTNM", "STABBR", "CONTROL", "DEGGRANT", "PSET4FLG")
CONTROL_LABELS = {"1": "Public", "2": "Private nonprofit", "3": "Private for-profit"}
CHANGE_REVIEW_FIELDS = (
    "unitid",
    "change_class",
    "change_reasons",
    "current_directory_present",
    "candidate_directory_present",
    "current_directory_eligible",
    "candidate_directory_eligible",
    "current_institution",
    "candidate_institution",
    "current_state",
    "candidate_state",
    "current_control_code",
    "current_control",
    "candidate_control_code",
    "candidate_control",
    "current_deggrant",
    "candidate_deggrant",
    "current_pset4flg",
    "candidate_pset4flg",
    "name_changed",
    "state_changed",
    "control_changed",
    "current_analysis_present",
    "current_nonresident_enrollment",
    "current_total_enrollment",
    "analytical_review_scope",
    "interpretation_boundary",
)
CHANGE_CLASS_ORDER = {
    "LEFT_DIRECTORY_ELIGIBLE": 0,
    "ENTERED_DIRECTORY_ELIGIBLE": 1,
    "SHARED_ATTRIBUTE_CHANGE": 2,
}
INTERPRETATION_BOUNDARY = (
    "Directory structural observation only; do not infer institutional opening/closure or 2025 enrollment movement."
)


def download(url: str) -> bytes:
    request = Request(url, headers={"User-Agent": USER_AGENT})
    try:
        with urlopen(request, timeout=90) as response:
            if response.status != 200:
                raise RuntimeError(f"HTTP {response.status} for {url}")
            return response.read()
    except (HTTPError, URLError) as exc:
        raise RuntimeError(f"Failed to download {url}: {exc}") from exc


def archive_member(payload: bytes, suffix: str) -> tuple[str, bytes]:
    try:
        archive = ZipFile(io.BytesIO(payload))
    except BadZipFile as exc:
        raise RuntimeError("Candidate directory response is not a valid ZIP archive") from exc
    members = [name for name in archive.namelist() if name.lower().endswith(suffix)]
    if len(members) != 1:
        raise RuntimeError(f"Expected exactly one {suffix} member; found {members}")
    name = members[0]
    return name, archive.read(name)


def decode_csv(payload: bytes) -> tuple[list[dict[str, str]], list[str], str]:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = payload.decode(encoding)
            reader = csv.DictReader(io.StringIO(text, newline=""))
            rows = list(reader)
            return rows, list(reader.fieldnames or []), encoding
        except UnicodeDecodeError:
            continue
    raise RuntimeError("Candidate directory CSV is not decodable as UTF-8 or cp1252")


def current_directory() -> tuple[list[dict[str, str]], list[str], str, str]:
    folder = RAW_ROOT / str(END_YEAR) / f"HD{END_YEAR}"
    revised = sorted(folder.glob("*_rv.csv"))
    paths = revised or sorted(folder.glob("*.csv"))
    if not paths:
        raise RuntimeError(f"Current-year directory CSV is missing from {folder}")
    path = paths[0]
    payload = path.read_bytes()
    rows, fields, encoding = decode_csv(payload)
    return rows, fields, encoding, path.name


def directory_index(rows: list[dict[str, str]], *, label: str) -> dict[str, dict[str, str]]:
    result: dict[str, dict[str, str]] = {}
    for row in rows:
        unitid = row["UNITID"].strip()
        if not unitid:
            raise RuntimeError(f"{label} directory contains a blank UNITID")
        if unitid in result:
            raise RuntimeError(f"{label} directory contains duplicate UNITID {unitid}")
        result[unitid] = row
    return result


def current_analysis_index() -> dict[str, dict[str, str]]:
    path = DATA / "processed" / f"institutions_{END_YEAR}.csv"
    if not path.is_file():
        raise RuntimeError(f"Current institution analysis is missing: {path.relative_to(ROOT)}")
    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        required = {"unitid", "institution", "state", "control", "total_enrollment", "nonresident_enrollment"}
        missing = required - set(reader.fieldnames or [])
        if missing:
            raise RuntimeError(f"Current institution analysis is missing fields: {sorted(missing)}")
        rows = list(reader)
    result: dict[str, dict[str, str]] = {}
    for row in rows:
        unitid = row["unitid"].strip()
        if unitid in result:
            raise RuntimeError(f"Current institution analysis contains duplicate UNITID {unitid}")
        result[unitid] = row
    return result


def dictionary_vars(payload: bytes) -> tuple[dict[str, str], int]:
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    names = {name.casefold(): name for name in workbook.sheetnames}
    if "varlist" not in names:
        raise RuntimeError(f"Candidate directory dictionary has no Varlist sheet: {workbook.sheetnames}")
    sheet = workbook[names["varlist"]]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    columns = {str(value).casefold(): index for index, value in enumerate(header)}
    for required in ("varname", "vartitle"):
        if required not in columns:
            raise RuntimeError(f"Candidate dictionary Varlist missing {required}")
    result: dict[str, str] = {}
    for row in rows:
        name = row[columns["varname"]]
        if name is None:
            continue
        title = row[columns["vartitle"]]
        result[str(name)] = "" if title is None else str(title)
    return result, len(result)


def eligible(rows: list[dict[str, str]]) -> dict[str, dict[str, str]]:
    result: dict[str, dict[str, str]] = {}
    for row in rows:
        if row["STABBR"].strip() not in CORE_STATES:
            continue
        if row["DEGGRANT"].strip() != "1" or row["PSET4FLG"].strip() not in {"1", "3"}:
            continue
        unitid = row["UNITID"].strip()
        if unitid in result:
            raise RuntimeError(f"Duplicate eligible UNITID {unitid}")
        result[unitid] = row
    return result


def control_counts(rows: dict[str, dict[str, str]]) -> dict[str, int]:
    counts = Counter(row["CONTROL"].strip() for row in rows.values())
    unexpected = sorted(set(counts) - set(CONTROL_LABELS))
    if unexpected:
        raise RuntimeError(f"Unexpected CONTROL values in candidate directory: {unexpected}")
    return {CONTROL_LABELS[code]: counts.get(code, 0) for code in CONTROL_LABELS}


def yes_no(value: bool) -> str:
    return "Yes" if value else "No"


def control_label(code: str) -> str:
    if not code:
        return ""
    if code not in CONTROL_LABELS:
        raise RuntimeError(f"Unexpected CONTROL value in directory change review: {code}")
    return CONTROL_LABELS[code]


def eligibility_reasons(
    before: dict[str, str] | None,
    after: dict[str, str] | None,
    *,
    direction: str,
) -> list[str]:
    if direction == "enter":
        if before is None:
            return ["new_directory_record"]
        reasons: list[str] = []
        if before["STABBR"].strip() not in CORE_STATES and after and after["STABBR"].strip() in CORE_STATES:
            reasons.append("state_entered_core_geography")
        if before["DEGGRANT"].strip() != "1" and after and after["DEGGRANT"].strip() == "1":
            reasons.append("degree_granting_status_gained")
        if before["PSET4FLG"].strip() not in {"1", "3"} and after and after["PSET4FLG"].strip() in {"1", "3"}:
            reasons.append("title_iv_first_look_status_gained")
    elif direction == "leave":
        if after is None:
            return ["directory_record_absent"]
        reasons = []
        if before and before["STABBR"].strip() in CORE_STATES and after["STABBR"].strip() not in CORE_STATES:
            reasons.append("state_left_core_geography")
        if before and before["DEGGRANT"].strip() == "1" and after["DEGGRANT"].strip() != "1":
            reasons.append("degree_granting_status_lost")
        if before and before["PSET4FLG"].strip() in {"1", "3"} and after["PSET4FLG"].strip() not in {"1", "3"}:
            reasons.append("title_iv_first_look_status_lost")
    else:
        raise RuntimeError(f"Unsupported eligibility direction: {direction}")
    if not reasons:
        raise RuntimeError(f"Could not explain directory eligibility {direction} transition")
    return reasons


def build_change_review_rows(
    current_rows: list[dict[str, str]],
    candidate_rows: list[dict[str, str]],
    current_eligible: dict[str, dict[str, str]],
    candidate_eligible: dict[str, dict[str, str]],
) -> list[dict[str, str]]:
    current_all = directory_index(current_rows, label="Current")
    candidate_all = directory_index(candidate_rows, label="Candidate")
    analysis = current_analysis_index()
    shared = set(current_eligible) & set(candidate_eligible)
    entering = set(candidate_eligible) - set(current_eligible)
    leaving = set(current_eligible) - set(candidate_eligible)
    rows: list[dict[str, str]] = []

    def make_row(unitid: str, change_class: str, reasons: list[str]) -> dict[str, str]:
        before = current_all.get(unitid)
        after = candidate_all.get(unitid)
        current_analysis = analysis.get(unitid)
        current_name = before["INSTNM"].strip() if before else ""
        candidate_name = after["INSTNM"].strip() if after else ""
        current_state = before["STABBR"].strip() if before else ""
        candidate_state = after["STABBR"].strip() if after else ""
        current_control_code = before["CONTROL"].strip() if before else ""
        candidate_control_code = after["CONTROL"].strip() if after else ""
        name_changed = bool(before and after and current_name != candidate_name)
        state_changed = bool(before and after and current_state != candidate_state)
        control_changed = bool(before and after and current_control_code != candidate_control_code)
        scopes: list[str] = []
        if change_class in {"ENTERED_DIRECTORY_ELIGIBLE", "LEFT_DIRECTORY_ELIGIBLE"}:
            scopes.append("Universe continuity")
        if control_changed:
            scopes.append("Control comparability")
        if state_changed:
            scopes.append("State comparability")
        if name_changed:
            scopes.append("Label continuity")
        if not scopes:
            scopes.append("Directory continuity")
        return {
            "unitid": unitid,
            "change_class": change_class,
            "change_reasons": ";".join(reasons),
            "current_directory_present": yes_no(before is not None),
            "candidate_directory_present": yes_no(after is not None),
            "current_directory_eligible": yes_no(unitid in current_eligible),
            "candidate_directory_eligible": yes_no(unitid in candidate_eligible),
            "current_institution": current_name,
            "candidate_institution": candidate_name,
            "current_state": current_state,
            "candidate_state": candidate_state,
            "current_control_code": current_control_code,
            "current_control": control_label(current_control_code),
            "candidate_control_code": candidate_control_code,
            "candidate_control": control_label(candidate_control_code),
            "current_deggrant": before["DEGGRANT"].strip() if before else "",
            "candidate_deggrant": after["DEGGRANT"].strip() if after else "",
            "current_pset4flg": before["PSET4FLG"].strip() if before else "",
            "candidate_pset4flg": after["PSET4FLG"].strip() if after else "",
            "name_changed": yes_no(name_changed),
            "state_changed": yes_no(state_changed),
            "control_changed": yes_no(control_changed),
            "current_analysis_present": yes_no(current_analysis is not None),
            "current_nonresident_enrollment": current_analysis["nonresident_enrollment"] if current_analysis else "",
            "current_total_enrollment": current_analysis["total_enrollment"] if current_analysis else "",
            "analytical_review_scope": ";".join(scopes),
            "interpretation_boundary": INTERPRETATION_BOUNDARY,
        }

    for unitid in leaving:
        rows.append(
            make_row(
                unitid,
                "LEFT_DIRECTORY_ELIGIBLE",
                eligibility_reasons(current_all.get(unitid), candidate_all.get(unitid), direction="leave"),
            )
        )
    for unitid in entering:
        rows.append(
            make_row(
                unitid,
                "ENTERED_DIRECTORY_ELIGIBLE",
                eligibility_reasons(current_all.get(unitid), candidate_all.get(unitid), direction="enter"),
            )
        )
    for unitid in shared:
        before = current_eligible[unitid]
        after = candidate_eligible[unitid]
        reasons = []
        if before["INSTNM"].strip() != after["INSTNM"].strip():
            reasons.append("institution_name_changed")
        if before["STABBR"].strip() != after["STABBR"].strip():
            reasons.append("state_changed_within_core")
        if before["CONTROL"].strip() != after["CONTROL"].strip():
            reasons.append("control_changed")
        if before["PSET4FLG"].strip() != after["PSET4FLG"].strip():
            reasons.append("title_iv_indicator_changed_within_eligible_values")
        if reasons:
            rows.append(make_row(unitid, "SHARED_ATTRIBUTE_CHANGE", reasons))

    rows.sort(key=lambda row: (CHANGE_CLASS_ORDER[row["change_class"]], int(row["unitid"])))
    if len({row["unitid"] for row in rows}) != len(rows):
        raise RuntimeError("Directory change review must contain one row per affected UNITID")
    return rows


def serialize_change_review(rows: list[dict[str, str]]) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.DictWriter(buffer, fieldnames=CHANGE_REVIEW_FIELDS, lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8")


def summarize_change_review(rows: list[dict[str, str]], csv_bytes: bytes) -> dict[str, object]:
    def count_class(value: str) -> int:
        return sum(row["change_class"] == value for row in rows)

    def sum_current(field: str, predicate) -> int:
        return sum(int(row[field] or 0) for row in rows if predicate(row))

    entry_reason_counts = Counter(
        reason
        for row in rows
        if row["change_class"] == "ENTERED_DIRECTORY_ELIGIBLE"
        for reason in row["change_reasons"].split(";")
    )
    exit_reason_counts = Counter(
        reason
        for row in rows
        if row["change_class"] == "LEFT_DIRECTORY_ELIGIBLE"
        for reason in row["change_reasons"].split(";")
    )
    return {
        "csv_path": str(DEFAULT_CHANGE_REVIEW_OUTPUT.relative_to(ROOT)).replace("\\", "/"),
        "csv_sha256": hashlib.sha256(csv_bytes).hexdigest(),
        "row_count": len(rows),
        "entered_directory_eligible": count_class("ENTERED_DIRECTORY_ELIGIBLE"),
        "left_directory_eligible": count_class("LEFT_DIRECTORY_ELIGIBLE"),
        "shared_attribute_change_rows": count_class("SHARED_ATTRIBUTE_CHANGE"),
        "entry_reason_counts": dict(sorted(entry_reason_counts.items())),
        "exit_reason_counts": dict(sorted(exit_reason_counts.items())),
        "current_analysis_present_leavers": sum(
            row["change_class"] == "LEFT_DIRECTORY_ELIGIBLE" and row["current_analysis_present"] == "Yes"
            for row in rows
        ),
        "current_nonresident_exposure_leavers": sum_current(
            "current_nonresident_enrollment", lambda row: row["change_class"] == "LEFT_DIRECTORY_ELIGIBLE"
        ),
        "current_total_enrollment_exposure_leavers": sum_current(
            "current_total_enrollment", lambda row: row["change_class"] == "LEFT_DIRECTORY_ELIGIBLE"
        ),
        "current_nonresident_exposure_control_changes": sum_current(
            "current_nonresident_enrollment",
            lambda row: row["change_class"] == "SHARED_ATTRIBUTE_CHANGE" and row["control_changed"] == "Yes",
        ),
        "current_nonresident_exposure_state_changes": sum_current(
            "current_nonresident_enrollment",
            lambda row: row["change_class"] == "SHARED_ATTRIBUTE_CHANGE" and row["state_changed"] == "Yes",
        ),
        "current_nonresident_exposure_name_changes": sum_current(
            "current_nonresident_enrollment",
            lambda row: row["change_class"] == "SHARED_ATTRIBUTE_CHANGE" and row["name_changed"] == "Yes",
        ),
    }


def load_release_snapshot() -> tuple[dict[str, object], dict[str, dict[str, object]]]:
    snapshot = json.loads(RELEASE_SNAPSHOT.read_text(encoding="utf-8"))
    if int(snapshot.get("current_analysis_end_year", -1)) != END_YEAR:
        raise RuntimeError("Release-monitor current year does not match project END_YEAR")
    candidate = END_YEAR + 1
    if int(snapshot.get("candidate_next_year", -1)) != candidate:
        raise RuntimeError("Release-monitor candidate year does not match END_YEAR + 1")
    files = snapshot.get("files")
    if not isinstance(files, list):
        raise RuntimeError("Release-monitor files payload is invalid")
    by_name = {str(row.get("file", "")): row for row in files}
    required_names = {f"HD{candidate}.zip", f"HD{candidate}_Dict.zip"}
    if not required_names.issubset(by_name):
        raise RuntimeError("Release-monitor snapshot is missing candidate directory assets")
    if not all(bool(by_name[name].get("available")) for name in required_names):
        raise RuntimeError("Candidate directory preflight requires both HD data and dictionary to be available")
    return snapshot, by_name


def build_capture() -> tuple[dict[str, object], list[dict[str, str]], bytes]:
    release_snapshot, release_files = load_release_snapshot()
    candidate = END_YEAR + 1
    data_name = f"HD{candidate}.zip"
    dictionary_name = f"HD{candidate}_Dict.zip"
    data_url = str(release_files[data_name]["url"])
    dictionary_url = str(release_files[dictionary_name]["url"])

    data_archive = download(data_url)
    dictionary_archive = download(dictionary_url)
    data_member, data_csv = archive_member(data_archive, ".csv")
    dictionary_member, dictionary_xlsx = archive_member(dictionary_archive, ".xlsx")

    candidate_rows, candidate_fields, candidate_encoding = decode_csv(data_csv)
    current_rows, current_fields, current_encoding, current_csv = current_directory()
    dictionary, dictionary_field_count = dictionary_vars(dictionary_xlsx)

    missing_data_fields = [field for field in REQUIRED_FIELDS if field not in candidate_fields]
    missing_dictionary_fields = [field for field in REQUIRED_FIELDS if field not in dictionary]
    if missing_data_fields or missing_dictionary_fields:
        raise RuntimeError(
            "Candidate directory required fields are incomplete: "
            f"data={missing_data_fields}; dictionary={missing_dictionary_fields}"
        )

    data_fields_missing_from_dictionary = sorted(set(candidate_fields) - set(dictionary))
    dictionary_variables_missing_from_data = sorted(set(dictionary) - set(candidate_fields))
    if data_fields_missing_from_dictionary or dictionary_variables_missing_from_data:
        raise RuntimeError(
            "Candidate directory data/dictionary field sets differ: "
            f"data_only={data_fields_missing_from_dictionary}; dictionary_only={dictionary_variables_missing_from_data}"
        )

    ids = [row["UNITID"].strip() for row in candidate_rows]
    duplicate_unitids = len(ids) - len(set(ids))
    if duplicate_unitids:
        raise RuntimeError(f"Candidate directory contains {duplicate_unitids} duplicate UNITIDs")

    candidate_states = {row["STABBR"].strip() for row in candidate_rows if row["STABBR"].strip() in CORE_STATES}
    current_states = {row["STABBR"].strip() for row in current_rows if row["STABBR"].strip() in CORE_STATES}
    if candidate_states != CORE_STATES:
        raise RuntimeError(f"Candidate directory core-state coverage mismatch: {sorted(CORE_STATES - candidate_states)}")
    if current_states != CORE_STATES:
        raise RuntimeError("Current directory no longer covers all 50 states plus DC")

    candidate_eligible = eligible(candidate_rows)
    current_eligible = eligible(current_rows)
    shared = set(candidate_eligible) & set(current_eligible)
    entering = set(candidate_eligible) - set(current_eligible)
    leaving = set(current_eligible) - set(candidate_eligible)

    candidate_controls = control_counts(candidate_eligible)
    current_controls = control_counts(current_eligible)
    control_changes = sum(
        current_eligible[unitid]["CONTROL"].strip() != candidate_eligible[unitid]["CONTROL"].strip()
        for unitid in shared
    )
    state_changes = sum(
        current_eligible[unitid]["STABBR"].strip() != candidate_eligible[unitid]["STABBR"].strip()
        for unitid in shared
    )
    name_changes = sum(
        current_eligible[unitid]["INSTNM"].strip() != candidate_eligible[unitid]["INSTNM"].strip()
        for unitid in shared
    )

    removed_fields = sorted(set(current_fields) - set(candidate_fields))
    added_fields = sorted(set(candidate_fields) - set(current_fields))
    required_titles = {field: dictionary[field] for field in REQUIRED_FIELDS}
    change_review_rows = build_change_review_rows(
        current_rows,
        candidate_rows,
        current_eligible,
        candidate_eligible,
    )
    change_review_bytes = serialize_change_review(change_review_rows)
    change_review_summary = summarize_change_review(change_review_rows, change_review_bytes)

    expected_change_rows = len(entering) + len(leaving) + sum(
        any(
            current_eligible[unitid][field].strip() != candidate_eligible[unitid][field].strip()
            for field in ("INSTNM", "STABBR", "CONTROL", "PSET4FLG")
        )
        for unitid in shared
    )
    if int(change_review_summary["row_count"]) != expected_change_rows:
        raise RuntimeError("Candidate directory change-review row count does not reconcile")
    if int(change_review_summary["entered_directory_eligible"]) != len(entering):
        raise RuntimeError("Candidate directory entry review count does not reconcile")
    if int(change_review_summary["left_directory_eligible"]) != len(leaving):
        raise RuntimeError("Candidate directory exit review count does not reconcile")

    payload: dict[str, object] = {
        "checked_at_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "current_analysis_end_year": END_YEAR,
        "candidate_next_year": candidate,
        "preflight_scope": "candidate institutional directory only; not Fall Enrollment",
        "status": "PASS",
        "year_advance_authorized": False,
        "authorization_boundary": (
            "Directory-side structural checks pass, but year advancement remains governed by the complete "
            "Fall Enrollment plus directory package and full schema review."
        ),
        "candidate_source": {
            "data_url": data_url,
            "data_archive_bytes": len(data_archive),
            "data_sha256": hashlib.sha256(data_archive).hexdigest(),
            "data_member": data_member,
            "data_encoding": candidate_encoding,
            "dictionary_url": dictionary_url,
            "dictionary_archive_bytes": len(dictionary_archive),
            "dictionary_sha256": hashlib.sha256(dictionary_archive).hexdigest(),
            "dictionary_member": dictionary_member,
        },
        "required_directory_fields": list(REQUIRED_FIELDS),
        "required_field_titles": required_titles,
        "missing_required_data_fields": missing_data_fields,
        "missing_required_dictionary_fields": missing_dictionary_fields,
        "candidate_directory_rows": len(candidate_rows),
        "candidate_field_count": len(candidate_fields),
        "candidate_dictionary_variable_count": dictionary_field_count,
        "candidate_full_field_dictionary_match": True,
        "candidate_fields_missing_from_dictionary": data_fields_missing_from_dictionary,
        "candidate_dictionary_variables_missing_from_data": dictionary_variables_missing_from_data,
        "candidate_core_states_dc_count": len(candidate_states),
        "candidate_duplicate_unitid_count": duplicate_unitids,
        "candidate_directory_eligible_rows": len(candidate_eligible),
        "candidate_control_counts": candidate_controls,
        "current_reference": {
            "year": END_YEAR,
            "directory_csv": current_csv,
            "encoding": current_encoding,
            "directory_rows": len(current_rows),
            "field_count": len(current_fields),
            "core_states_dc_count": len(current_states),
            "directory_eligible_rows": len(current_eligible),
            "control_counts": current_controls,
        },
        "schema_delta": {
            "field_count_change": len(candidate_fields) - len(current_fields),
            "removed_fields": removed_fields,
            "added_fields": added_fields,
            "required_fields_preserved": not missing_data_fields and not missing_dictionary_fields,
        },
        "eligible_universe_continuity": {
            "shared_unitids": len(shared),
            "entering_unitids": len(entering),
            "leaving_unitids": len(leaving),
            "net_change": len(candidate_eligible) - len(current_eligible),
            "shared_control_changes": control_changes,
            "shared_state_changes": state_changes,
            "shared_name_changes": name_changes,
        },
        "change_review": change_review_summary,
        "release_monitor_checked_at_utc": release_snapshot["checked_at_utc"],
        "release_monitor_decision": release_snapshot["release_decision"],
    }
    return payload, change_review_rows, change_review_bytes


def build_payload() -> dict[str, object]:
    payload, _, _ = build_capture()
    return payload


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--write-snapshot",
        type=Path,
        default=None,
        help=f"Optional output path. Recommended: {DEFAULT_OUTPUT.relative_to(ROOT)}",
    )
    parser.add_argument(
        "--write-change-review",
        type=Path,
        default=None,
        help=(
            "Optional CSV output for institution-level directory changes. When --write-snapshot is used "
            f"without this option, defaults to {DEFAULT_CHANGE_REVIEW_OUTPUT.relative_to(ROOT)}."
        ),
    )
    args = parser.parse_args()

    try:
        payload, change_review_rows, change_review_bytes = build_capture()
    except RuntimeError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1

    print(json.dumps(payload, indent=2))
    change_review_path = args.write_change_review
    if args.write_snapshot and change_review_path is None:
        change_review_path = DEFAULT_CHANGE_REVIEW_OUTPUT
    if change_review_path:
        change_review_path.parent.mkdir(parents=True, exist_ok=True)
        change_review_path.write_bytes(change_review_bytes)
        print(
            f"Wrote candidate directory change review ({len(change_review_rows)} rows) to {change_review_path}",
            file=sys.stderr,
        )
    if args.write_snapshot:
        args.write_snapshot.parent.mkdir(parents=True, exist_ok=True)
        args.write_snapshot.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")
        print(f"Wrote candidate directory preflight to {args.write_snapshot}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
